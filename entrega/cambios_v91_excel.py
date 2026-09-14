"""Genera el cuadro de cambios entre el v9 y el v9.1, en Excel.

Organizado en las TRES categorías con que el cliente lee una versión, y no por
área técnica:

    1. Detección automática   lo que el sistema encuentra y calcula solo
    2. Usabilidad y errores   lo que se arregló o dejó de estorbar
    3. Nuevas funcionalidades herramientas y pantallas que antes no existían

Una hoja por categoría para poder copiar y pegar por bloques, una hoja con todo
junto para filtrar, y una con el antes/después medido de los errores.

    uv run --with openpyxl python entrega/cambios_v91_excel.py

Fuente: `entrega/V9.1.md` (lo que entra) y las correcciones P23–P30 de
`debug/PENDIENTES.md`.
"""
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

SALIDA = Path("entrega/Cambios_v9_a_v9.1.xlsx")

DETECCION = "1 · Detección automática"
ERRORES = "2 · Usabilidad y errores"
NUEVAS = "3 · Nuevas funcionalidades"

CATEGORIAS = [
    (DETECCION,
     "Lo que el sistema encuentra y calcula solo: el proceso de detección y las "
     "cifras que salen de él.",
     "1F3B57"),
    (ERRORES,
     "Lo que se arregló o dejó de estorbar. Mucho de esto no se ve, pero es lo "
     "que hace que el trabajo no se pierda.",
     "8A5A0B"),
    (NUEVAS,
     "Herramientas y pantallas que antes no existían.",
     "1D6F42"),
]

# Categoría | Subárea | Tipo | Cambio | Qué hace y por qué importa | Atajo | Origen
FILAS = [
    # ======================= 1 · DETECCIÓN AUTOMÁTICA =======================
    (DETECCION, "Proceso de detección", "Mejora", "Corte de ruido interactivo",
     "El proceso se parte en dos y se detiene a la mitad: el usuario elige sobre la máscara cuánto ruido cortar antes de que siga. Es una pantalla nueva en el asistente.",
     "", "Del equipo"),
    (DETECCION, "Proceso de detección", "Mejora", "Nodos nuevos de detección",
     "Vista previa del percentil, intensidad de ambos signos, detección de «fuera de vista» por velocidad y fusión de trayectorias paralelas.",
     "", "Del equipo"),
    (DETECCION, "Proceso de detección", "Corrección", "Miniaturas del detector",
     "Arreglo de la generación de miniaturas en el detector de tronadura.",
     "", "Del equipo"),
    (DETECCION, "Cálculo y métricas", "Corrección", "El histograma medía otra cosa",
     "El botón decía «histograma de alcances» pero graficaba cuánto se salió la roca del borde del área de voladura, que es otra magnitud. Ahora grafica el alcance desde el tiro de origen: la misma cifra que entrega el archivo de resultados.",
     "", "Nuestro"),
    (DETECCION, "Cálculo y métricas", "Corrección", "El diámetro de la voladura se inventaba",
     "Si el análisis no traía la zona de origen, el sistema asumía una voladura de 1,13 m y todas las trayectorias caían en la primera barra del histograma, sin un solo aviso. Ahora se dice que falta el dato en vez de inventarlo.",
     "", "Nuestro"),
    (DETECCION, "Cálculo y métricas", "Corrección", "La clase se recalcula al editar",
     "Recortar, unir, dibujar o alisar cambia dónde termina el vuelo, y la clase —peligrosa, proyección o fuera de vista— se heredaba sin volver a mirar: una recortada seguía diciendo «fuera de vista» aunque el recorte le quitara justo la punta por la que lo era.",
     "", "Nuestro"),
    (DETECCION, "Cálculo y métricas", "Corrección", "El escape del área se recalcula al editar",
     "Una trayectoria dibujada a mano nacía con escape cero fijo y salía en el archivo de resultados como «no salió del área», que es falso. Ahora se recalcula al dibujar, al recortar y al unir.",
     "", "Nuestro"),
    (DETECCION, "Cálculo y métricas", "Corrección", "La pantalla y el archivo contaban distinto",
     "El archivo afirmaba un pozo de origen para trayectorias que la pantalla se negaba a asociar: 176 sin empalme en pantalla contra 75 en el archivo, del mismo caso. Ahora los dos aplican la misma regla.",
     "", "Nuestro"),

    # ==================== 2 · USABILIDAD Y ERRORES ==========================
    (ERRORES, "Guardar y retomar", "Mejora", "El avance se guarda en el análisis",
     "Hasta el v9, «Guardar avance» producía un archivo y nada más: retomar dependía de que ese archivo siguiera en la carpeta de Descargas. Ahora el trabajo vive dentro del propio análisis.",
     "", "Nuestro"),
    (ERRORES, "Guardar y retomar", "Mejora", "Reabrir devuelve la pantalla como se dejó",
     "Abrir un análisis con avance lo aplica solo: vuelven los umbrales, las capas, las clases apagadas, el modo de trabajo, el alisado y las aprobadas.",
     "", "Nuestro"),
    (ERRORES, "Guardar y retomar", "Mejora", "Ver el análisis original",
     "Vuelve al resultado crudo del proceso sin borrar el avance guardado.",
     "", "Nuestro"),
    (ERRORES, "Guardar y retomar", "Mejora", "Abrir un archivo de avance al entrar",
     "Un archivo de avance soltado en la pantalla de entrada va directo a la vista final, en su propio análisis. Sirve para traer trabajo de otro equipo o para tener varias alternativas del mismo caso.",
     "", "Nuestro"),
    (ERRORES, "Guardar y retomar", "Corrección", "El avance pertenece a su pasada",
     "Rebobinar y recalcular con otro corte de ruido borra el avance: al recalcular, los identificadores de roca se reasignan y aplicar el avance anterior sería incorrecto en silencio.",
     "", "Nuestro"),
    (ERRORES, "Interno (no se ve)", "Mejora", "El trabajo vive en la base del sistema",
     "El avance pasa a guardarse junto al análisis, con un resumen leído aparte para que la lista no tenga que cargar megas de puntos cada vez que se abre.",
     "", "Nuestro"),
    (ERRORES, "Interno (no se ve)", "Mejora", "Borrado real de un análisis",
     "Borra el registro y su carpeta de archivos, con la ruta verificada antes de eliminar nada.",
     "", "Nuestro"),
    (ERRORES, "Interno (no se ve)", "Mejora", "Purga por antigüedad, apagada por defecto",
     "Se puede activar por configuración. Viene apagada a propósito: ahí vive el trabajo del cliente, y borrárselo sin que lo pida es peor que quedarse sin disco.",
     "", "Nuestro"),
    (ERRORES, "Interno (no se ve)", "Mejora", "El sistema sabe en qué estado está cada análisis",
     "Informa si un análisis se puede abrir, por qué no, cuánto pesa y en qué paso retomarlo. Es lo que permite que la lista ofrezca lo correcto.",
     "", "Nuestro"),
    (ERRORES, "Errores corregidos", "Corrección", "Se abría la pantalla equivocada",
     "Al agregarse el paso del corte de ruido, la vista de limpieza se corrió un número y la lista seguía abriendo la anterior. Abría otra pantalla sin dar ningún error.",
     "", "Nuestro"),
    (ERRORES, "Errores corregidos", "Corrección", "Análisis pausado, invisible e imborrable",
     "Un análisis que quedó esperando el corte de ruido no aparecía en la lista y tampoco se podía borrar. Cerrar el navegador en ese paso lo perdía para siempre.",
     "", "Nuestro"),
    (ERRORES, "Errores corregidos", "Corrección", "Un lazo se llevaba lo aprobado",
     "El lazo de descarte y el recorte arrasaban también con las trayectorias ya aprobadas. Ahora las respetan y lo dicen; para sacar una aprobada hay que seleccionarla y descartarla a mano.",
     "", "Nuestro"),
    (ERRORES, "Errores corregidos", "Corrección", "El contador de aprobadas no bajaba",
     "Contaba las aprobadas sin mirar si seguían en el trabajo, así que el número nunca bajaba. El caso extremo daba un archivo de resultados con cero trayectorias activas y veinte aprobadas.",
     "", "Nuestro"),
    (ERRORES, "Errores corregidos", "Corrección", "Ocultar por clase no respondía",
     "Ocultar por clase funcionaba en un video nuevo, pero no sobre un trabajo cargado: lo aprobado ignoraba la casilla. Es un control de vista, no un juicio de calidad, y ahora alcanza a todo.",
     "", "Nuestro"),
    (ERRORES, "Errores corregidos", "Corrección", "Aviso antes de exportar sin origen",
     "Si hay trayectorias sin tiro de origen, el histograma las mide por el trazo visible —mucho más corto que el vuelo real— y ahora avisa antes de bajar la imagen.",
     "", "Nuestro"),
    (ERRORES, "Errores corregidos", "Corrección", "La exportación ya no falla sin radio",
     "Un caso sin radio de evacuación rompía la generación del histograma entero, sin ningún mensaje.",
     "", "Nuestro"),
    (ERRORES, "Errores corregidos", "Corrección", "Textos ilegibles en la pantalla de entrada",
     "La pantalla usaba colores de tema claro sobre el fondo azul de la aplicación: el pie de página y los textos de apoyo quedaban negro sobre azul.",
     "", "Nuestro"),
    (ERRORES, "Errores corregidos", "Corrección", "No había vuelta desde la otra vista",
     "Al pasarse a la vista anterior no quedaba ningún botón para volver: había que recargar la página, y eso perdía la limpieza no guardada.",
     "", "Nuestro"),
    (ERRORES, "Errores corregidos", "Corrección", "Cambiar de vista perdía el trabajo",
     "Cambiar de vista cerraba la vista de limpieza entera, con todo lo trabajado desde el último guardado. Ahora las dos quedan vivas y solo se alterna cuál se ve.",
     "", "Nuestro"),
    (ERRORES, "Errores corregidos", "Corrección", "Cambiar de vista podía sacarte del asistente",
     "Si el análisis se había abierto desde la lista, el botón de la vista anterior devolvía al paso 1 y se perdía el trabajo. Ahora se explica por qué no está disponible, en vez de llevar a un callejón.",
     "", "Nuestro"),
    (ERRORES, "Interno (no se ve)", "Corrección", "El armador borraba el paquete de otra versión",
     "Con un nombre de versión con punto, el script de empaquetado tomaba el «.1» por extensión y su limpieza borraba el paquete de OTRA versión ya entregada.",
     "", "Nuestro"),

    # ================= 3 · NUEVAS FUNCIONALIDADES (UX) ======================
    (NUEVAS, "Revisión de trayectorias", "Mejora", "Aprobar una trayectoria",
     "Marca la trayectoria como buena y la saca del lienzo. Lo aprobado ya no lo tocan los filtros ni un lazo: sin esto, tres horas de revisión se deshacían con un umbral mal movido.",
     "Enter", "Nuestro"),
    (NUEVAS, "Revisión de trayectorias", "Mejora", "Cosechar",
     "Descarta de una vez todo lo que quedó sin aprobar. Es el cierre del trabajo: lo revisado se queda, el resto se va, y avisa cuántas hay de cada cosa antes de hacerlo.",
     "", "Nuestro"),
    (NUEVAS, "Revisión de trayectorias", "Mejora", "Color por avance",
     "El trazo responde a una sola pregunta mientras se trabaja: ¿ya pasé por esta o no? El color por clase —la conclusión— queda para el entregable, y las exportaciones lo fuerzan pase lo que pase.",
     "P", "Nuestro"),
    (NUEVAS, "Revisión de trayectorias", "Mejora", "Lupa de trabajo",
     "Ver todo, solo el lienzo (lo que falta), solo las aprobadas (lo cosechado) o solo las que quedaron sin tiro de origen. No cambia el estado de nada: es qué se dibuja.",
     "", "Nuestro"),
    (NUEVAS, "Revisión de trayectorias", "Mejora", "Aislar una trayectoria",
     "Una trayectoria sola en pantalla, con las otras 3.700 apagadas.",
     "I", "Nuestro"),
    (NUEVAS, "Revisión de trayectorias", "Mejora", "Alisar y enderezar",
     "El zigzag fino de una traza no es física: es el centro del objeto saltando unos píxeles entre cuadros. Se reajusta a la curva que la física permite, con dos pasadas distintas —una para el ruido, otra para los puntos saltados— y fuerza regulable.",
     "S · Shift+S", "Nuestro"),
    (NUEVAS, "Revisión de trayectorias", "Mejora", "Dar vuelta el sentido de un trazo",
     "Una trayectoria dibujada al revés buscaba su pozo en la dirección contraria y se quedaba sin origen, sin forma de arreglarlo salvo borrarla y volver a dibujarla.",
     "V", "Nuestro"),
    (NUEVAS, "Revisión de trayectorias", "Mejora", "Cuántas quedaron sin empalme, y cuáles",
     "El panel dice cuántas trayectorias quedaron sin tiro de origen y por qué —pocos puntos, sin tangente, ningún pozo detrás, o descartadas por el calce temporal— y una lupa las junta en pantalla para corregirlas.",
     "", "Nuestro"),
    (NUEVAS, "Pantalla de entrada", "Mejora", "Lista de análisis (pantalla nueva)",
     "El ingreso ya no cae en el paso 1: cae en la lista de lo que hay en el equipo, con video, fecha, trayectorias, peso en disco y si tiene avance guardado.",
     "", "Nuestro"),
    (NUEVAS, "Pantalla de entrada", "Mejora", "Reabrir no reprocesa",
     "Abrir un análisis ya hecho no vuelve a correr el proceso: son minutos de cálculo que no se pagan dos veces.",
     "", "Nuestro"),
    (NUEVAS, "Pantalla de entrada", "Mejora", "Abrir donde quedó",
     "El terminado va a la edición; el que quedó esperando el corte de ruido vuelve a su control, con la máscara ya calculada.",
     "", "Nuestro"),
    (NUEVAS, "Pantalla de entrada", "Mejora", "Borrar de a uno",
     "Elimina un análisis y sus archivos, avisando si tiene aprobadas guardadas. Antes la única forma de liberar disco era «CERRAR Y BORRAR TODO», que se llevaba también lo que se quería conservar.",
     "", "Nuestro"),
    (NUEVAS, "Pantalla de entrada", "Mejora", "No se ofrece lo que no sirve",
     "Los análisis de versiones anteriores, los que fallaron y los que están procesando se ven, pero dicen por qué no se abren.",
     "", "Nuestro"),
    (NUEVAS, "Asistente", "Mejora", "Ayudas de calibración y reporte",
     "Guía de matriz, cronómetro, máxima proyección general y dentro del área de seguridad, mover un punto de la proyección, empalme y filtro por puntos, invertir la asociación al pozo.",
     "", "Del equipo"),
    (NUEVAS, "Asistente", "Mejora", "La otra vista se abre con lo limpiado",
     "Al cambiar a la vista anterior se le pasa lo que llevas limpiado, en vez de la maraña original. Al volver se avisa de que lo editado allá no se trae: los dos modelos no son compatibles y mezclarlos daría un resultado incorrecto.",
     "", "Nuestro"),
]

EVIDENCIA = [
    (DETECCION, "Histograma: trayectorias en la primera barra",
     "726 de 726 (máximo medido: 3,86 m)",
     "repartidas en 21 barras, hasta 206 m"),
    (DETECCION, "Recortar al 25 % cuatro «fuera de vista»",
     "las 4 seguían diciendo «fuera de vista»",
     "3 «proyección» y 1 «proyección peligrosa»"),
    (DETECCION, "Pantalla contra archivo: trayectorias sin empalme",
     "176 contra 75, del mismo caso", "176 en los dos"),
    (ERRORES, "Aprobadas que sobreviven a un lazo de descarte",
     "0 de 20", "20 de 20"),
    (ERRORES, "Archivo de resultados después de ese lazo",
     "0 trayectorias activas y 20 aprobadas (imposible)",
     "20 activas y 20 aprobadas"),
    (ERRORES, "Ocultar por clase en un trabajo con 200 aprobadas",
     "quedaban 141 visibles de la clase apagada", "0"),
    (ERRORES, "Contraste del texto del pie en la pantalla de entrada",
     "1,07 : 1 (ilegible)", "7,6 : 1"),
    (NUEVAS, "Trayectorias sin tiro de origen",
     "un solo motivo, y era el equivocado",
     "176, desglosadas en 4 motivos reales"),
    (NUEVAS, "Curva dibujada al revés",
     "sin pozo de origen y sin arreglo posible",
     "se da vuelta con la tecla V y recupera su tiro"),
]

GRIS = "F2F4F6"
BORDE = Border(bottom=Side(style="thin", color="D8DEE4"))
COLOR_CAT = {c: col for c, _, col in CATEGORIAS}


def encabezar(ws, titulos, anchos, color="1F3B57"):
    ws.append(titulos)
    for i, (t, a) in enumerate(zip(titulos, anchos), start=1):
        c = ws.cell(row=1, column=i)
        c.font = Font(bold=True, color="FFFFFF", size=11)
        c.fill = PatternFill("solid", fgColor=color)
        c.alignment = Alignment(vertical="center", horizontal="left")
        ws.column_dimensions[get_column_letter(i)].width = a
    ws.row_dimensions[1].height = 24
    ws.freeze_panes = "A2"


def pintar(ws, fila, ncols, par, ancho_cols, tipo=None, col_tipo=None):
    for col in range(1, ncols + 1):
        c = ws.cell(row=fila, column=col)
        c.alignment = Alignment(vertical="top", wrap_text=(col in ancho_cols))
        c.border = BORDE
        if par:
            c.fill = PatternFill("solid", fgColor=GRIS)
    if tipo and col_tipo:
        ws.cell(row=fila, column=col_tipo).font = Font(
            color="B45309" if tipo == "Corrección" else "1D6F42", bold=True)


wb = Workbook()

# ----------------------------------------------------------------- Resumen
ws = wb.active
ws.title = "Resumen"
encabezar(ws, ["Categoría", "Qué agrupa", "Cambios", "Mejoras", "Correcciones"],
          [30, 74, 11, 11, 14])
for nombre, desc, color in CATEGORIAS:
    de_esta = [f for f in FILAS if f[0] == nombre]
    mej = sum(1 for f in de_esta if f[2] == "Mejora")
    ws.append([nombre, desc, len(de_esta), mej, len(de_esta) - mej])
    f = ws.max_row
    for col in range(1, 6):
        c = ws.cell(row=f, column=col)
        c.alignment = Alignment(vertical="center", wrap_text=(col == 2))
        c.border = BORDE
    ws.cell(row=f, column=1).font = Font(bold=True, color=color, size=12)
    ws.row_dimensions[f].height = 46
ws.append(["TOTAL", "", len(FILAS),
           sum(1 for f in FILAS if f[2] == "Mejora"),
           sum(1 for f in FILAS if f[2] == "Corrección")])
for col in range(1, 6):
    ws.cell(row=ws.max_row, column=col).font = Font(bold=True)

ws.append([])
ws.append(["v9.1 sobre el v9. El v9 trajo la edición de trayectorias; el v9.1 hace que el trabajo deje de perderse."])
ws.cell(row=ws.max_row, column=1).font = Font(italic=True, color="6B7280")

# -------------------------------------------------- Una hoja por categoría
NOMBRE_HOJA = {DETECCION: "1 Detección automática",
               ERRORES: "2 Usabilidad y errores",
               NUEVAS: "3 Nuevas funcionalidades"}

for nombre, desc, color in CATEGORIAS:
    hoja = wb.create_sheet(NOMBRE_HOJA[nombre])
    encabezar(hoja, ["#", "Grupo", "Tipo", "Cambio", "Qué hace y por qué importa",
                     "Atajo", "Origen"],
              [5, 24, 13, 40, 96, 12, 12], color)
    for i, f in enumerate([x for x in FILAS if x[0] == nombre], start=1):
        _, sub, tipo, cambio, detalle, atajo, origen = f
        hoja.append([i, sub, tipo, cambio, detalle, atajo, origen])
        r = hoja.max_row
        pintar(hoja, r, 7, i % 2 == 0, (4, 5), tipo, 3)
        hoja.cell(row=r, column=4).font = Font(bold=True)
        hoja.row_dimensions[r].height = 48
    hoja.auto_filter.ref = f"A1:G{hoja.max_row}"

# --------------------------------------------------------- Todo, filtrable
wt = wb.create_sheet("Todo (filtrable)")
encabezar(wt, ["#", "Categoría", "Grupo", "Tipo", "Cambio",
               "Qué hace y por qué importa", "Atajo", "Origen"],
          [5, 26, 24, 13, 40, 92, 12, 12])
for i, f in enumerate(FILAS, start=1):
    cat, sub, tipo, cambio, detalle, atajo, origen = f
    wt.append([i, cat, sub, tipo, cambio, detalle, atajo, origen])
    r = wt.max_row
    pintar(wt, r, 8, i % 2 == 0, (5, 6), tipo, 4)
    wt.cell(row=r, column=2).font = Font(bold=True, color=COLOR_CAT[cat])
    wt.cell(row=r, column=5).font = Font(bold=True)
    wt.row_dimensions[r].height = 48
wt.auto_filter.ref = f"A1:H{wt.max_row}"

# ------------------------------------------------------ Correcciones medidas
wm = wb.create_sheet("Correcciones medidas")
encabezar(wm, ["Categoría", "Qué se midió", "Antes", "Ahora"], [26, 52, 48, 44])
for i, (cat, a, b, c) in enumerate(EVIDENCIA, start=1):
    wm.append([cat, a, b, c])
    r = wm.max_row
    pintar(wm, r, 4, i % 2 == 0, (1, 2, 3, 4))
    wm.cell(row=r, column=1).font = Font(bold=True, color=COLOR_CAT[cat], size=10)
    wm.cell(row=r, column=2).font = Font(bold=True)
    wm.cell(row=r, column=3).font = Font(color="B45309")
    wm.cell(row=r, column=4).font = Font(color="1D6F42")
    wm.row_dimensions[r].height = 34
wm.append([])
wm.append(["Medido ejecutando la vista sobre el caso 3160-789, de 726 trayectorias."])
wm.cell(row=wm.max_row, column=1).font = Font(italic=True, color="6B7280")

SALIDA.parent.mkdir(parents=True, exist_ok=True)
wb.save(SALIDA)

print(f"{SALIDA}")
for nombre, _, _ in CATEGORIAS:
    de_esta = [f for f in FILAS if f[0] == nombre]
    mej = sum(1 for f in de_esta if f[2] == "Mejora")
    print(f"  {nombre:28} {len(de_esta):>3} cambios  ({mej} mejoras, {len(de_esta)-mej} correcciones)")
print(f"  {'TOTAL':28} {len(FILAS):>3}")
